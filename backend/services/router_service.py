import os
import re
import csv
import json
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Dict, Any, List, Tuple, Optional

# Pipeline imports
try:
    from pipelines import ingest_text, ingest_image, ingest_audio, ingest_video, ingest_link
except ImportError:
    ingest_text = None
    ingest_image = None
    ingest_audio = None
    ingest_video = None
    ingest_link = None

from backend.config import settings

# 46 Supported File Formats categorized by media pipeline
DOC_EXTS = {".pdf", ".docx", ".doc", ".pptx", ".ppt", ".xlsx", ".xls", ".csv", ".tsv", ".txt", ".log", ".md", ".markdown", ".rtf"}
WEB_EXTS = {".xml", ".rss", ".atom"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".svg", ".tiff", ".tif", ".bmp"}
AUDIO_EXTS = {".mp3", ".wav", ".m4a", ".ogg", ".flac"}
VIDEO_EXTS = {".mp4", ".mkv", ".mov", ".avi", ".webm"}
CYBER_EXTS = {".stix", ".taxii", ".json", ".jsonl", ".evtx", ".syslog", ".yara", ".sigma", ".py", ".sh", ".ps1"}

ALL_SUPPORTED_EXTENSIONS = (
    DOC_EXTS | WEB_EXTS | IMAGE_EXTS | AUDIO_EXTS | VIDEO_EXTS | CYBER_EXTS
)

def extract_regex_iocs(text: str) -> Dict[str, List[str]]:
    """Extracts critical cybersecurity indicators without data loss."""
    cves = list(set(re.findall(r"\bCVE-\d{4}-\d{4,7}\b", text, re.IGNORECASE)))
    ips = list(set(re.findall(r"\b(?:\d{1,3}\.){3}\d{1,3}\b", text)))
    sha256s = list(set(re.findall(r"\b[a-fA-F0-9]{64}\b", text)))
    domains = list(set(re.findall(r"\b(?:[a-zA-Z0-9-]+\.)+(?:com|org|net|io|in|gov|internal|local)\b", text, re.IGNORECASE)))
    return {
        "cves": cves,
        "ips": [ip for ip in ips if not ip.startswith("0.") and not ip.startswith("255.")],
        "sha256s": sha256s,
        "domains": domains,
    }


class PipelineRouterService:
    @staticmethod
    def detect_media_pipeline(filename: str) -> str:
        ext = Path(filename).suffix.lower()
        if ext in IMAGE_EXTS:
            return "image"
        if ext in AUDIO_EXTS:
            return "audio"
        if ext in VIDEO_EXTS:
            return "video"
        return "text"

    @staticmethod
    def process_file_into_context(file_path: str, filename: str, session_id: str = "") -> Dict[str, Any]:
        ext = Path(filename).suffix.lower()
        pipeline_type = PipelineRouterService.detect_media_pipeline(filename)

        markdown = ""
        citations = []
        structured_metadata = {"filename": filename, "extension": ext, "pipeline": pipeline_type}

        try:
            # 1. Image Pipeline Routing (.png, .jpg, .jpeg, .webp, .svg, .tiff, .tif, .bmp)
            if pipeline_type == "image":
                media_url = f"/api/pipeline/media/{session_id}/{filename}" if session_id else f"/api/pipeline/media/default/{filename}"
                structured_metadata["media_url"] = media_url

                if ext == ".svg":
                    try:
                        svg_text = Path(file_path).read_text(encoding="utf-8", errors="replace")
                        markdown = f"### SVG Vector Graphics Diagram: {filename}\n```xml\n{svg_text[:2000]}\n```"
                        citations.append({
                            "id": f"img-{filename}",
                            "kind": "ocr",
                            "label": f"[Vector Graphic] {filename}",
                            "media_url": media_url,
                            "bbox": {"x": 5.0, "y": 5.0, "width": 90.0, "height": 90.0},
                        })
                    except Exception:
                        pass
                if not markdown and ingest_image:
                    r = ingest_image(file_path)
                    markdown = r.markdown_output or ""
                    all_boxes = getattr(r, "all_boxes", []) or []

                    for anchor in getattr(r, "grounding_sources", []):
                        bbox = getattr(anchor, "bbox", None)
                        bbox_dict = bbox.model_dump() if hasattr(bbox, "model_dump") else (bbox if isinstance(bbox, dict) else None)
                        anchor_id = getattr(anchor, "id", f"img-{len(citations)+1}")
                        clean_anchor_id = anchor_id.replace("^", "")
                        citations.append({
                            "id": clean_anchor_id,
                            "kind": "ocr",
                            "label": f"[{getattr(anchor, 'visual_anchor', filename)}] {getattr(anchor, 'extracted_verbatim', '')[:100]}",
                            "bbox": bbox_dict,
                            "media_url": media_url,
                            "all_boxes": all_boxes,
                        })

                    structured_metadata["grounding_sources"] = [
                        a.model_dump() if hasattr(a, "model_dump") else str(a)
                        for a in getattr(r, "grounding_sources", [])
                    ]
                    structured_metadata["all_boxes"] = all_boxes

            # 2. Audio Pipeline Routing (.mp3, .wav, .m4a, .ogg, .flac)
            elif pipeline_type == "audio":
                if ingest_audio:
                    r = ingest_audio(file_path)
                    markdown = r.markdown_output or ""
                    for anchor in getattr(r, "grounding_sources", []):
                        citations.append({
                            "id": getattr(anchor, "id", f"aud-{len(citations)+1}"),
                            "kind": "file",
                            "label": f"[{getattr(anchor, 'temporal_anchor', filename)}] {getattr(anchor, 'extracted_verbatim', '')[:100]}",
                        })
                    structured_metadata["grounding_sources"] = [a.model_dump() if hasattr(a, "model_dump") else str(a) for a in getattr(r, "grounding_sources", [])]

            # 3. Video Pipeline Routing (.mp4, .mkv, .mov, .avi, .webm)
            elif pipeline_type == "video":
                if ingest_video:
                    r = ingest_video(file_path, save_outputs=False, enrich=False)
                    markdown = r.clean_markdown or ""
                    for scene in getattr(r, "scenes", []):
                        citations.append({
                            "id": f"vid-scene-{getattr(scene, 'scene_id', len(citations)+1)}",
                            "kind": "file",
                            "label": f"[{getattr(scene, 'timestamp_display', '00:00')}] {getattr(scene, 'spoken_transcript', '')[:80]}",
                        })
                    structured_metadata["scenes_count"] = len(getattr(r, "scenes", []))

            # 4. Text & Structured Data Pipeline Routing
            else:
                markdown, citations, extra_meta = PipelineRouterService._process_text_and_cyber_formats(file_path, filename, ext, session_id=session_id)
                structured_metadata.update(extra_meta)

        except Exception as e:
            # Resilient fallback extraction
            raw_content = ""
            try:
                raw_content = Path(file_path).read_text(encoding="utf-8", errors="replace")[:4000]
            except Exception:
                pass
            markdown = f"## Ingested Source File: {filename} ({ext})\n\n{raw_content or f'Extracted telemetry & intelligence from {filename}.'}"
            citations.append({"id": f"src-{filename}", "kind": "file", "label": filename})

        # Pre-LLM deterministic indicator extraction across all extracted text
        iocs = extract_regex_iocs(markdown)
        for cve in iocs["cves"]:
            citations.append({"id": f"cve-{cve}", "kind": "text", "label": f"Vulnerability: {cve}"})
        for ip in iocs["ips"][:5]:
            citations.append({"id": f"ip-{ip}", "kind": "text", "label": f"Host/IP: {ip}"})

        structured_metadata["iocs"] = iocs
        return {
            "markdown": markdown,
            "citations": citations,
            "metadata": structured_metadata,
            "pipeline": pipeline_type,
        }

    @staticmethod
    def _process_text_and_cyber_formats(file_path: str, filename: str, ext: str, session_id: str = "") -> Tuple[str, List[Dict[str, Any]], Dict[str, Any]]:
        markdown = ""
        citations = []
        meta = {}

        # Office Presentations (.pptx, .ppt)
        if ext in (".pptx", ".ppt"):
            try:
                from pptx import Presentation
                prs = Presentation(file_path)
                slides_md = []
                for idx, slide in enumerate(prs.slides, 1):
                    slide_texts = []
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            for paragraph in shape.text_frame.paragraphs:
                                text = paragraph.text.strip()
                                if text:
                                    slide_texts.append(text)
                    slide_body = "\n".join(f"- {t}" for t in slide_texts)
                    slides_md.append(f"#### Slide {idx}\n{slide_body}")
                markdown = f"### Presentation Deck: {filename}\n\n" + "\n\n".join(slides_md)
                citations.append({"id": f"ppt-{filename}", "kind": "file", "label": f"Presentation {filename} ({len(prs.slides)} slides)"})
                meta["slide_count"] = len(prs.slides)
                return markdown, citations, meta
            except Exception:
                pass

        # Spreadsheets (.xlsx, .xls, .csv, .tsv)
        if ext in (".xlsx", ".xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, data_only=True)
                sheets_md = []
                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    rows = list(ws.iter_rows(values_only=True))
                    if rows:
                        header = [str(cell or "") for cell in rows[0]]
                        table_lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * len(header)) + " |"]
                        for row in rows[1:50]:  # Cap to first 50 rows
                            table_lines.append("| " + " | ".join(str(cell or "") for cell in row) + " |")
                        sheets_md.append(f"#### Sheet: {sheet_name}\n" + "\n".join(table_lines))
                markdown = f"### Spreadsheet Data: {filename}\n\n" + "\n\n".join(sheets_md)
                citations.append({"id": f"sheet-{filename}", "kind": "file", "label": f"Spreadsheet {filename}"})
                return markdown, citations, meta
            except Exception:
                pass

        if ext in (".csv", ".tsv"):
            try:
                delimiter = "\t" if ext == ".tsv" else ","
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    reader = csv.reader(f, delimiter=delimiter)
                    rows = list(reader)
                    if rows:
                        header = rows[0]
                        lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * len(header)) + " |"]
                        for row in rows[1:50]:
                            lines.append("| " + " | ".join(row) + " |")
                        markdown = f"### Delimited Telemetry Table: {filename}\n\n" + "\n".join(lines)
                        citations.append({"id": f"table-{filename}", "kind": "file", "label": f"Data Table {filename}"})
                        return markdown, citations, meta
            except Exception:
                pass

        # Web & Syndication (.xml, .rss, .atom)
        if ext in (".xml", ".rss", ".atom"):
            try:
                tree = ET.parse(file_path)
                root = tree.getroot()
                items_md = []
                for elem in root.iter():
                    if elem.tag.endswith(("item", "entry")):
                        title = elem.findtext(".//title") or elem.findtext("title") or "Syndicated Alert"
                        link = elem.findtext(".//link") or elem.findtext("link") or ""
                        summary = elem.findtext(".//description") or elem.findtext(".//summary") or ""
                        items_md.append(f"- **{title}**\n  - Link: {link}\n  - Summary: {summary[:300]}")
                if items_md:
                    markdown = f"### Syndicated Feed: {filename}\n\n" + "\n\n".join(items_md)
                else:
                    xml_str = ET.tostring(root, encoding="utf-8").decode("utf-8", errors="replace")[:3000]
                    markdown = f"### XML Document: {filename}\n```xml\n{xml_str}\n```"
                citations.append({"id": f"feed-{filename}", "kind": "file", "label": f"Web Feed {filename}"})
                return markdown, citations, meta
            except Exception:
                pass

        # Cybersecurity Structured Data (.stix, .taxii, .json, .jsonl)
        if ext in (".stix", ".taxii", ".json", ".jsonl"):
            try:
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    if ext == ".jsonl":
                        records = [json.loads(line) for line in f if line.strip()]
                        markdown = f"### Structured Event Stream (JSONL): {filename} ({len(records)} events)\n```json\n{json.dumps(records[:10], indent=2)}\n```"
                    else:
                        obj = json.load(f)
                        if isinstance(obj, dict) and obj.get("type") == "bundle":
                            # STIX 2.1 Bundle
                            stix_objects = obj.get("objects", [])
                            indicators = [o.get("name") or o.get("pattern") for o in stix_objects if o.get("type") == "indicator"]
                            markdown = f"### STIX 2.1 Cyber Threat Intelligence Bundle: {filename}\n- Total STIX Objects: {len(stix_objects)}\n- Key Indicators: {', '.join(str(i) for i in indicators[:10])}\n\n```json\n{json.dumps(stix_objects[:5], indent=2)}\n```"
                        else:
                            markdown = f"### Structured Data: {filename}\n```json\n{json.dumps(obj, indent=2)[:3000]}\n```"
                citations.append({"id": f"cyber-{filename}", "kind": "file", "label": f"Cybersecurity Dataset: {filename}"})
                return markdown, citations, meta
            except Exception:
                pass

        # Threat Detection Rules & Security Scripts (.yara, .sigma, .py, .sh, .ps1, .syslog, .evtx)
        if ext in (".yara", ".sigma", ".py", ".sh", ".ps1", ".syslog", ".evtx"):
            try:
                content = Path(file_path).read_text(encoding="utf-8", errors="replace")
                syntax = "yaml" if ext == ".sigma" else "python" if ext == ".py" else "bash" if ext == ".sh" else "powershell" if ext == ".ps1" else "text"
                markdown = f"### Security Rule & Script Artifact: {filename}\n```{syntax}\n{content[:4000]}\n```"
                citations.append({"id": f"rule-{filename}", "kind": "file", "label": f"Detection Logic: {filename}"})
                return markdown, citations, meta
            except Exception:
                pass

        # Standard Document Parser (PDF, DOCX, TXT, MD, RTF, LOG) via Text Ingestion Pipeline
        if ingest_text and ext in (".pdf", ".docx", ".doc", ".txt", ".log", ".md", ".markdown", ".rtf"):
            try:
                r = ingest_text(file_path, save_outputs=False, enrich=False)
                markdown = r.clean_markdown or ""
                citations.append({"id": f"doc-{filename}", "kind": "file", "label": filename})
                if hasattr(r, "iocs") and r.iocs:
                    for cve in getattr(r.iocs, "cves", []):
                        citations.append({"id": f"cve-{cve}", "kind": "file", "label": f"Vulnerability: {cve}"})
                    for ip in getattr(r.iocs, "ipv4_addresses", []):
                        citations.append({"id": f"ip-{ip}", "kind": "file", "label": f"IOC IP: {ip}"})

                # Visual PDF Parsing: Render pages with visual diagrams/scans and extract coordinates
                if ext == ".pdf":
                    try:
                        import pymupdf
                        from pipelines.image_pipeline.extractors.ocr_utils import ImageOCRUtils
                        doc = pymupdf.open(file_path)
                        sess_token = session_id or "default"
                        renders_dir = settings.RENDERS_DIR / sess_token
                        renders_dir.mkdir(parents=True, exist_ok=True)

                        for p_idx in range(len(doc)):
                            page = doc[p_idx]
                            raw_p_text = page.get_text().strip()
                            page_imgs = page.get_images()

                            # Hybrid Detection: If page has embedded diagrams or low selectable text (< 120 chars)
                            if len(page_imgs) > 0 or len(raw_p_text) < 120:
                                safe_stem = re.sub(r"[^\w\-]", "_", Path(filename).stem)
                                render_filename = f"{safe_stem}_p{p_idx+1}.png"
                                render_file = renders_dir / render_filename
                                pix = page.get_pixmap(dpi=150)
                                pix.save(str(render_file))

                                # Run OCR coordinate detection on the rasterized page
                                ocr_data = ImageOCRUtils.detect_ocr_boxes(str(render_file))
                                p_boxes = ocr_data.get("boxes", [])
                                rel_url = f"/api/pipeline/media/{sess_token}/renders/{render_filename}"

                                top_bbox = p_boxes[0]["bbox"] if p_boxes else {"x": 5.0, "y": 5.0, "width": 90.0, "height": 90.0}
                                top_text = p_boxes[0]["text"] if p_boxes else f"Visual Page {p_idx+1}"

                                citations.append({
                                    "id": f"pdf-vis-p{p_idx+1}",
                                    "kind": "ocr",
                                    "label": f"[Page {p_idx+1} Visual Diagram] {top_text[:75]}",
                                    "bbox": top_bbox,
                                    "media_url": rel_url,
                                    "page_number": p_idx + 1,
                                    "all_boxes": p_boxes,
                                })
                        doc.close()
                    except Exception as pdf_err:
                        print(f"[RouterService PDF Visual Parse Notice]: {pdf_err}")

                return markdown, citations, meta
            except Exception:
                pass

        # Fallback pure text read
        content = Path(file_path).read_text(encoding="utf-8", errors="replace")
        markdown = f"### Source Document: {filename}\n\n{content}"
        citations.append({"id": f"doc-{filename}", "kind": "file", "label": filename})
        return markdown, citations, meta

    @staticmethod
    def process_link_into_context(url: str, output_dir: Optional[str] = None) -> Dict[str, Any]:
        """
        Executes link_pipeline to scrape a web document and extract grounded context.
        """
        if ingest_link:
            r = ingest_link(url, output_dir=output_dir, save_outputs=True)
            citations = []
            if getattr(r, "grounding_sources", None):
                for anchor in r.grounding_sources:
                    citations.append({
                        "id": getattr(anchor, "id", f"link-{len(citations)+1}"),
                        "kind": "link",
                        "label": f"[{r.metadata.domain}] {getattr(anchor, 'extracted_verbatim', '')[:100]}",
                    })
            else:
                citations.append({
                    "id": f"link-{r.metadata.domain}",
                    "kind": "link",
                    "label": f"[{r.metadata.domain}] {r.metadata.title}",
                })

            return {
                "markdown": r.clean_markdown or "",
                "citations": citations,
                "metadata": r.metadata.model_dump(),
                "iocs": r.iocs.model_dump() if hasattr(r, "iocs") else {},
                "md_file_path": r.md_file_path,
                "json_file_path": r.json_file_path,
                "pipeline": "link",
                "success": r.success,
                "error_message": r.error_message,
            }
        return {
            "markdown": f"### Web Link Source\nURL: {url}",
            "citations": [{"id": f"link-{url[:30]}", "kind": "link", "label": url}],
            "metadata": {"url": url},
            "iocs": {},
            "md_file_path": None,
            "json_file_path": None,
            "pipeline": "link",
            "success": False,
            "error_message": "link_pipeline module not loaded",
        }

router_service = PipelineRouterService()
