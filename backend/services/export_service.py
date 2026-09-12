"""
Transmute Deliverable Export Service.
Provides high-fidelity conversion of final cybersecurity deliverables into:
- Markdown (.md)
- Plain Text (.txt)
- PDF Document (.pdf)
- Microsoft Word Document (.docx)
- Multi-deliverable ZIP Archives (.zip)

CRITICAL REQUIREMENT:
All citation markers (e.g., [^src-1], [^1], [^aud-2], HTML citation pills)
are completely stripped from exported files. Citations are intended solely
for in-app provenance exploration and must not appear in customer deliverables.
"""

from __future__ import annotations

import io
import re
import zipfile
import logging
from typing import Any, Dict, List, Tuple
from datetime import datetime

import markdown
import pymupdf
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

logger = logging.getLogger("transmute.export_service")

DELIVERABLE_TITLES: Dict[str, str] = {
    "linkedin_post": "LinkedIn Executive Post",
    "social_thread": "Social Media Thread",
    "advisory": "Technical Security Advisory",
    "exec_summary": "Executive Intelligence Summary",
    "incident_report": "Cybersecurity Incident Report",
    "press_release": "Public Security Statement",
    "slide_deck": "Executive Briefing Slide Deck",
    "video_script": "Audio & Video Briefing Script",
    "playbook": "Remediation & Incident Playbook",
}


def strip_citations(text: str) -> str:
    """
    Strips all markdown citation anchors, footnotes, and HTML citation pills.
    Cleans up any dangling spaces before punctuation.
    """
    if not text:
        return ""
    # Strip HTML citation pills: <button ... class="...citation-pill...">...</button>
    clean = re.sub(
        r'<button[^>]*class=["\'][^"\']*citation-pill[^"\']*["\'][^>]*>.*?</button>',
        "",
        text,
        flags=re.DOTALL | re.IGNORECASE,
    )
    # Strip markdown footnote citations: [^src-1], [^1], [^img-2], [^aud-1], etc.
    clean = re.sub(r'\[\^[^\]]+\]', '', clean)
    # Strip lingering inline html tags
    clean = re.sub(r'</?(?:span|div|button|p|a)[^>]*>', '', clean, flags=re.IGNORECASE)
    # Correct spaces before punctuation (e.g., "attack ." -> "attack.")
    clean = re.sub(r'\s+([,.:;!?])', r'\1', clean)
    # Collapse multiple inline spaces to single space while preserving line breaks
    clean = re.sub(r'[ \t]{2,}', ' ', clean)
    return clean.strip()


def markdown_to_plain_text(markdown_text: str) -> str:
    """
    Converts markdown content to readable, clean plain text without citations.
    """
    clean = strip_citations(markdown_text)
    # Remove horizontal rules
    clean = re.sub(r'^\s*[-*_]{3,}\s*$', '', clean, flags=re.MULTILINE)
    # Format headings nicely with underline
    def format_heading(match: re.Match) -> str:
        level = len(match.group(1))
        title = match.group(2).strip()
        underline = "=" * min(len(title), 50) if level == 1 else "-" * min(len(title), 40)
        return f"\n{title}\n{underline}\n"

    clean = re.sub(r'^(#{1,6})\s*(.+)$', format_heading, clean, flags=re.MULTILINE)
    # Bold / italic markers
    clean = re.sub(r'\*\*(.+?)\*\*', r'\1', clean)
    clean = re.sub(r'__(.+?)__', r'\1', clean)
    clean = re.sub(r'\*(.+?)\*', r'\1', clean)
    clean = re.sub(r'_(.+?)_', r'\1', clean)
    # Inline code
    clean = re.sub(r'`(.+?)`', r'\1', clean)
    # Markdown links: [text](url) -> text (url)
    clean = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', clean)
    # Blockquotes
    clean = re.sub(r'^\s*>\s?', '  | ', clean, flags=re.MULTILINE)
    # Normalize excessive newlines
    clean = re.sub(r'\n{3,}', '\n\n', clean)
    return clean.strip()


def markdown_to_pdf_bytes(markdown_text: str, output_type: str = "") -> bytes:
    """
    Converts markdown content into a PDF document using PyMuPDF Story.
    Ensures citations are cleanly removed and the document has executive styling.
    """
    clean_md = strip_citations(markdown_text)
    title_text = DELIVERABLE_TITLES.get(output_type, output_type.replace("_", " ").title()) or "Security Deliverable"
    current_date = datetime.now().strftime("%B %d, %Y")

    # Convert markdown to clean HTML
    html_body = markdown.markdown(
        clean_md,
        extensions=["extra", "tables", "nl2br", "sane_lists"],
    )

    styled_html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  @page {{
    margin: 40pt 45pt;
  }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
    font-size: 10.5pt;
    line-height: 1.55;
    color: #1e293b;
    margin: 0;
    padding: 0;
  }}
  .header-banner {{
    border-bottom: 2px solid #0284c7;
    padding-bottom: 12pt;
    margin-bottom: 18pt;
  }}
  .meta-tag {{
    display: inline-block;
    background-color: #f0fdf4;
    color: #166534;
    font-size: 8.5pt;
    font-weight: 700;
    padding: 2pt 8pt;
    border-radius: 4pt;
    text-transform: uppercase;
    letter-spacing: 0.5pt;
    border: 1px solid #bbf7d0;
  }}
  .date-label {{
    float: right;
    font-size: 9pt;
    color: #64748b;
    font-weight: 500;
  }}
  h1 {{
    color: #0f172a;
    font-size: 18pt;
    font-weight: 700;
    margin-top: 8pt;
    margin-bottom: 6pt;
    line-height: 1.25;
  }}
  h2 {{
    color: #0369a1;
    font-size: 13pt;
    font-weight: 600;
    margin-top: 16pt;
    margin-bottom: 6pt;
    border-bottom: 1px solid #e2e8f0;
    padding-bottom: 3pt;
  }}
  h3 {{
    color: #334155;
    font-size: 11.5pt;
    font-weight: 600;
    margin-top: 12pt;
    margin-bottom: 4pt;
  }}
  p {{
    margin-top: 0;
    margin-bottom: 8pt;
    color: #334155;
  }}
  ul, ol {{
    margin-top: 2pt;
    margin-bottom: 8pt;
    padding-left: 20pt;
  }}
  li {{
    margin-bottom: 4pt;
    color: #334155;
  }}
  blockquote {{
    border-left: 3.5pt solid #0284c7;
    background-color: #f8fafc;
    margin: 8pt 0;
    padding: 6pt 12pt;
    color: #475569;
    font-style: italic;
    border-radius: 0 4pt 4pt 0;
  }}
  table {{
    border-collapse: collapse;
    width: 100%;
    margin: 10pt 0;
    font-size: 9.5pt;
  }}
  th, td {{
    border: 1px solid #cbd5e1;
    padding: 6pt 8pt;
    text-align: left;
  }}
  th {{
    background-color: #f1f5f9;
    color: #0f172a;
    font-weight: 600;
  }}
  tr:nth-child(even) {{
    background-color: #f8fafc;
  }}
  code {{
    background-color: #f1f5f9;
    color: #0f172a;
    padding: 1pt 4pt;
    border-radius: 3pt;
    font-family: monospace;
    font-size: 9pt;
  }}
  pre {{
    background-color: #0f172a;
    color: #f8fafc;
    padding: 8pt 10pt;
    border-radius: 4pt;
    overflow-x: auto;
    font-size: 8.5pt;
    line-height: 1.4;
  }}
  pre code {{
    background-color: transparent;
    color: inherit;
    padding: 0;
  }}
  .footer-note {{
    margin-top: 24pt;
    border-top: 1px solid #e2e8f0;
    padding-top: 8pt;
    font-size: 8pt;
    color: #94a3b8;
    text-align: center;
  }}
</style>
</head>
<body>
  <div class="header-banner">
    <span class="meta-tag">Validated Security Deliverable</span>
    <span class="date-label">{current_date}</span>
    <h1>{title_text}</h1>
  </div>
  {html_body}
  <div class="footer-note">
    Generated by SIH Cyber Intelligence Platform • Confidential & Proprietary
  </div>
</body>
</html>
"""

    try:
        story = pymupdf.Story(html=styled_html)

        def rect_fn(rect_num, filled):
            mediabox = pymupdf.Rect(0, 0, 612, 792)  # Standard Letter
            rect = pymupdf.Rect(45, 40, 567, 752)    # Margins
            return mediabox, rect, None

        doc = story.write_with_links(rect_fn)
        pdf_bytes = doc.tobytes()
        doc.close()
        return pdf_bytes
    except Exception as exc:
        logger.warning(f"PyMuPDF Story render failed ({exc}), falling back to direct textbox pagination")
        # Direct fallback
        doc = pymupdf.open()
        plain_text = markdown_to_plain_text(clean_md)
        page = doc.new_page(width=612, height=792)
        rect = pymupdf.Rect(45, 45, 567, 747)
        page.insert_text((45, 40), title_text, fontsize=16, fontname="helv", color=(0.01, 0.52, 0.78))
        page.insert_textbox(rect, plain_text, fontsize=10.5, fontname="helv")
        pdf_bytes = doc.tobytes()
        doc.close()
        return pdf_bytes


def _set_cell_background(cell, fill_hex: str):
    """Sets the background XML shading of a docx table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)


def markdown_to_docx_bytes(markdown_text: str, output_type: str = "") -> bytes:
    """
    Converts markdown content into a Microsoft Word (.docx) document using python-docx.
    Strips citations and applies typography, tables, bullet points, and headings.
    """
    clean_md = strip_citations(markdown_text)
    title_text = DELIVERABLE_TITLES.get(output_type, output_type.replace("_", " ").title()) or "Security Deliverable"

    doc = docx.Document()

    # Set 0.8 inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(title_text)
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(2, 132, 199) # #0284c7

    # Subtitle meta
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(14)
    meta_run = meta_p.add_run(f"Confidential Security Deliverable  |  Generated on {datetime.now().strftime('%B %d, %Y')}")
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = RGBColor(100, 116, 139)

    lines = clean_md.split("\n")
    i = 0
    in_code_block = False
    code_lines: List[str] = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Code blocks
        if stripped.startswith("```"):
            if in_code_block:
                # End code block
                code_p = doc.add_paragraph()
                code_p.paragraph_format.left_indent = Inches(0.25)
                code_run = code_p.add_run("\n".join(code_lines))
                code_run.font.name = "Consolas"
                code_run.font.size = Pt(9)
                code_run.font.color.rgb = RGBColor(30, 41, 59)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Blank lines
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if re.match(r'^[-*_]{3,}$', stripped):
            i += 1
            continue

        # Headings
        if stripped.startswith("### "):
            h = doc.add_heading(stripped[4:], level=3)
            if h.runs:
                h.runs[0].font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue
        if stripped.startswith("## "):
            h = doc.add_heading(stripped[3:], level=2)
            if h.runs:
                h.runs[0].font.color.rgb = RGBColor(3, 105, 161)
            i += 1
            continue
        if stripped.startswith("# "):
            h = doc.add_heading(stripped[2:], level=1)
            if h.runs:
                h.runs[0].font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue

        # Markdown Table Detection
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines: List[str] = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Process table
            parsed_rows = []
            for t_line in table_lines:
                # Ignore separator row (| :--- | :--- |)
                if re.match(r'^\|(?:\s*:?-+:?\s*\|)+$', t_line):
                    continue
                cells = [c.strip() for c in t_line.strip("|").split("|")]
                parsed_rows.append(cells)

            if parsed_rows:
                num_cols = max(len(r) for r in parsed_rows)
                docx_table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                docx_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                docx_table.style = "Table Grid"

                for row_idx, row_data in enumerate(parsed_rows):
                    for col_idx in range(num_cols):
                        cell = docx_table.cell(row_idx, col_idx)
                        text_val = row_data[col_idx] if col_idx < len(row_data) else ""
                        cell.text = text_val
                        if row_idx == 0:
                            _set_cell_background(cell, "F1F5F9")
                            for run in cell.paragraphs[0].runs:
                                run.font.bold = True
                                run.font.size = Pt(9.5)
                        else:
                            if row_idx % 2 == 0:
                                _set_cell_background(cell, "F8FAFC")
                            for run in cell.paragraphs[0].runs:
                                run.font.size = Pt(9)
                doc.add_paragraph() # Spacing
            continue

        # Bullet List Items
        if re.match(r'^[-*+]\s+', stripped):
            item_text = re.sub(r'^[-*+]\s+', '', stripped)
            p = doc.add_paragraph(style='List Bullet')
            _add_formatted_runs_to_paragraph(p, item_text)
            i += 1
            continue

        # Numbered List Items
        if re.match(r'^\d+\.\s+', stripped):
            item_text = re.sub(r'^\d+\.\s+', '', stripped)
            p = doc.add_paragraph(style='List Number')
            _add_formatted_runs_to_paragraph(p, item_text)
            i += 1
            continue

        # Blockquote
        if stripped.startswith(">"):
            quote_text = re.sub(r'^>\s*', '', stripped)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(quote_text)
            run.font.italic = True
            run.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        # Standard Paragraph with bold/italics
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        _add_formatted_runs_to_paragraph(p, stripped)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_formatted_runs_to_paragraph(paragraph, text: str):
    """
    Parses **bold** and *italic* and inline `code` into distinct docx Runs.
    """
    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.font.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.italic = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(15, 23, 42)
        else:
            paragraph.add_run(token)


def export_deliverable(content: str, output_type: str, file_format: str) -> Tuple[bytes, str, str]:
    """
    Exports a single deliverable into the requested file format.
    Supported formats: 'md', 'txt', 'pdf', 'docx'.
    Returns: (file_bytes, filename, media_type)
    """
    fmt = file_format.lower().lstrip(".")
    base_slug = output_type.lower().replace(" ", "_") or "deliverable"

    if fmt == "md":
        clean = strip_citations(content)
        data = clean.encode("utf-8")
        filename = f"{base_slug}.md"
        media_type = "text/markdown; charset=utf-8"
    elif fmt == "txt":
        plain = markdown_to_plain_text(content)
        data = plain.encode("utf-8")
        filename = f"{base_slug}.txt"
        media_type = "text/plain; charset=utf-8"
    elif fmt == "pdf":
        data = markdown_to_pdf_bytes(content, output_type=output_type)
        filename = f"{base_slug}.pdf"
        media_type = "application/pdf"
    elif fmt == "docx":
        data = markdown_to_docx_bytes(content, output_type=output_type)
        filename = f"{base_slug}.docx"
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        raise ValueError(f"Unsupported export format: {file_format}. Supported: md, txt, pdf, docx")

    return data, filename, media_type


def export_all_deliverables_zip(deliverables: List[Dict[str, Any]], file_format: str, prefix: str = "cyber_deliverables") -> Tuple[bytes, str, str]:
    """
    Exports all supplied deliverables converted into the requested format (md, txt, pdf, docx)
    and packages them inside a single .zip archive.
    Returns: (zip_bytes, zip_filename, "application/zip")
    """
    fmt = file_format.lower().lstrip(".")
    if fmt not in ("md", "txt", "pdf", "docx"):
        raise ValueError(f"Unsupported zip export format: {file_format}. Supported: md, txt, pdf, docx")

    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        used_names: set[str] = set()

        for idx, item in enumerate(deliverables):
            output_type = item.get("output_type") or item.get("outputType") or f"deliverable_{idx+1}"
            content = item.get("content") or item.get("text") or ""

            data, item_filename, _ = export_deliverable(content, output_type, fmt)

            # Avoid collision in zip
            final_filename = item_filename
            counter = 1
            while final_filename in used_names:
                name_part, ext = item_filename.rsplit(".", 1)
                final_filename = f"{name_part}_{counter}.{ext}"
                counter += 1

            used_names.add(final_filename)
            zf.writestr(final_filename, data)

    zip_filename = f"{prefix}_{fmt}.zip"
    return zip_buf.getvalue(), zip_filename, "application/zip"
